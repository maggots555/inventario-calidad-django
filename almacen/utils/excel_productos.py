"""
Excel del catálogo de productos activos.

EXPLICACIÓN PARA PRINCIPIANTES:
-------------------------------
La vista HTTP solo pide el archivo. Aquí se arma el libro: una hoja,
encabezados y una fila por producto. No sabe de permisos ni de URLs.
"""

from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# Azul de la marca SIC. Openpyxl pide el color sin el símbolo #.
_AZUL_SIC = '1F6391'

# (encabezado, ancho de columna). El orden es el de las columnas A, B, C…
_COLUMNAS = (
    ('Código', 18),
    ('Nombre', 36),
    ('Descripción', 40),
    ('Tipo', 28),
    ('Categoría', 22),
    ('Sucursal', 22),
    ('Ubicación física', 18),
    ('Stock actual', 14),
    ('Stock mínimo', 14),
    ('Stock máximo', 14),
    ('Estado de stock', 16),
    ('Clave SAT', 14),
    ('Costo unitario', 16),
    ('Valor del stock', 16),
    ('Proveedor', 28),
    ('Días de reposición', 18),
    ('Creado', 20),
    ('Última actualización', 22),
)

# Columnas numéricas (1-based) para que Excel las trate como número, no texto.
_COLUMNAS_ENTERAS = (8, 9, 10, 16)
_COLUMNAS_DINERO = (13, 14)
_COLUMNAS_FECHA = (17, 18)


def construir_excel_productos_activos(productos):
    """
    Genera el .xlsx de productos activos listos para descargar.

    Objetivo de negocio:
        Quien está en la pantalla de productos se lleva el detalle del
        catálogo (stock, costos, clave SAT, sucursal) en una hoja.

    Args:
        productos: iterable de ProductoAlmacen. Quien llama ya aplicó
            activo=True y los filtros de la pantalla. Aquí no se filtra.

    Returns:
        bytes del archivo .xlsx. Si no hay productos, el archivo igual
        trae la fila de encabezados.

    Efectos secundarios:
        Ninguno sobre la base de datos. Solo lee atributos de cada producto.
    """
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Productos activos'

    # Paso 1: encabezado fijo. La fila 1 siempre son los nombres de columna.
    _escribir_encabezado(hoja)

    # Paso 2: una fila por producto, en el orden en que llegó el iterable.
    for numero_fila, producto in enumerate(productos, start=2):
        _escribir_fila(hoja, numero_fila, producto)

    # El filtro de Excel abarca encabezado y datos para poder ordenar en el archivo.
    hoja.freeze_panes = 'A2'
    ultima_columna = get_column_letter(len(_COLUMNAS))
    hoja.auto_filter.ref = f'A1:{ultima_columna}{hoja.max_row}'

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _escribir_encabezado(hoja):
    """
    Pinta la fila 1 con el azul SIC y deja el ancho de cada columna.

    Args:
        hoja: hoja activa de openpyxl.

    Efectos secundarios:
        Modifica la hoja en memoria. No toca disco ni base de datos.
    """
    relleno = PatternFill(start_color=_AZUL_SIC, end_color=_AZUL_SIC, fill_type='solid')
    letra = Font(bold=True, color='FFFFFF', size=11)
    centrado = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for indice, (titulo, ancho) in enumerate(_COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=indice, value=titulo)
        celda.fill = relleno
        celda.font = letra
        celda.alignment = centrado
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    hoja.row_dimensions[1].height = 22


def _escribir_fila(hoja, numero_fila, producto):
    """
    Escribe un producto en la fila indicada, con formato de dinero y fecha.

    Args:
        hoja: hoja de openpyxl.
        numero_fila: número de fila (2 es el primer producto).
        producto: ProductoAlmacen.

    Efectos secundarios:
        Modifica la hoja en memoria.
    """
    # get_estado_stock devuelve (código, nombre, clase CSS). Al Excel va el nombre.
    _estado_codigo, estado_nombre, _clase_css = producto.get_estado_stock()
    valores = (
        producto.codigo_producto,
        producto.nombre,
        producto.descripcion or '',
        producto.get_tipo_producto_display(),
        producto.categoria.nombre if producto.categoria_id else '',
        producto.sucursal.nombre if producto.sucursal_id else 'Almacén central',
        producto.ubicacion_fisica or '',
        producto.stock_actual,
        producto.stock_minimo,
        producto.stock_maximo,
        estado_nombre,
        producto.clave_sat or '',
        producto.costo_unitario,
        producto.valor_total_stock(),
        producto.proveedor_principal.nombre if producto.proveedor_principal_id else '',
        producto.tiempo_reposicion_dias,
        _fecha_para_excel(producto.fecha_creacion),
        _fecha_para_excel(producto.fecha_actualizacion),
    )

    for indice, valor in enumerate(valores, start=1):
        celda = hoja.cell(row=numero_fila, column=indice, value=valor)
        if indice in _COLUMNAS_ENTERAS:
            celda.number_format = '0'
        elif indice in _COLUMNAS_DINERO:
            celda.number_format = '#,##0.00'
        elif indice in _COLUMNAS_FECHA and valor is not None:
            celda.number_format = 'DD/MM/YYYY HH:MM'


def _fecha_para_excel(fecha):
    """
    Convierte una fecha de Django a una fecha ingenua en hora local.

    Objetivo de negocio:
        Excel no guarda zona horaria. La hora que ve el usuario es la
        del país activo, no UTC.

    Args:
        fecha: datetime (consciente o ingenuo) o None.

    Returns:
        datetime sin tzinfo, o None si no hay fecha.

    Efectos secundarios:
        Ninguno.
    """
    if fecha is None:
        return None
    if timezone.is_aware(fecha):
        fecha = timezone.localtime(fecha)
    return fecha.replace(tzinfo=None)
