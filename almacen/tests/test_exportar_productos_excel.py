"""
Descarga Excel de los productos activos del almacén.

Objetivo de negocio:
    El botón de la lista baja el mismo catálogo que se está viendo,
    con el detalle de stock y clave SAT, y deja fuera a los inactivos.
"""

from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Permission, User
from django.test import RequestFactory, SimpleTestCase, TestCase, override_settings
from django.urls import resolve, reverse
from openpyxl import load_workbook

from almacen import views as almacen_views
from almacen.models import ProductoAlmacen
from almacen.views_catalogo import lista_productos
from almacen.views_productos_excel import exportar_productos_excel


STORAGES_TEST = {
    'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
    'staticfiles': {
        'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage',
    },
}

_CONTENT_TYPE_XLSX = (
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
)


class ExportarProductosExcelHumoTest(SimpleTestCase):
    """La URL del Excel apunta a la vista nueva, no al monolito."""

    def test_url_resuelve_a_la_vista_nueva(self):
        """reverse/resolve y el reexport de views.py son la misma función."""
        coincidencia = resolve(reverse('almacen:exportar_productos_excel'))
        self.assertIs(coincidencia.func, exportar_productos_excel)
        self.assertIs(
            almacen_views.exportar_productos_excel,
            exportar_productos_excel,
        )


@override_settings(STORAGES=STORAGES_TEST)
class ExportarProductosExcelTest(TestCase):
    """Contenido del archivo y permiso para descargarlo."""

    def setUp(self):
        """
        Un usuario con permiso de ver productos y tres artículos.

        Efectos secundarios:
            Inserta usuarios y productos en la base de prueba.
        """
        self.usuario = User.objects.create_user(
            username='almacen_excel',
            password='excel-test',
        )
        permiso = Permission.objects.get(codename='view_productoalmacen')
        self.usuario.user_permissions.add(permiso)
        self.sin_permiso = User.objects.create_user(
            username='almacen_sin_excel',
            password='excel-test',
        )
        self.factory = RequestFactory()
        ProductoAlmacen.objects.create(
            codigo_producto='CON-SAT',
            nombre='Con clave',
            clave_sat='43211600',
            stock_actual=2,
            costo_unitario=Decimal('10.50'),
            descripcion='Pasta térmica',
        )
        ProductoAlmacen.objects.create(
            codigo_producto='SIN-SAT',
            nombre='Sin clave',
            clave_sat='',
            stock_actual=0,
        )
        ProductoAlmacen.objects.create(
            codigo_producto='INACTIVO',
            nombre='Oculto',
            clave_sat='11111111',
            activo=False,
        )

    def _descargar(self, usuario, parametros=None):
        """
        Llama la vista de Excel sin el middleware de país.

        Args:
            usuario: usuario autenticado de la petición.
            parametros: dict de GET. None equivale a la lista sin filtros.

        Returns:
            HttpResponse de exportar_productos_excel.
        """
        request = self.factory.get('/almacen/productos/excel/', parametros or {})
        request.user = usuario
        return exportar_productos_excel(request)

    def _codigos(self, respuesta):
        """
        Lee la columna Código del archivo descargado.

        Args:
            respuesta: HttpResponse con el .xlsx.

        Returns:
            list[str]: códigos desde la fila 2.
        """
        libro = load_workbook(BytesIO(respuesta.content))
        hoja = libro.active
        return [
            hoja.cell(row=fila, column=1).value
            for fila in range(2, hoja.max_row + 1)
        ]

    def test_incluye_activos_y_deja_fuera_al_inactivo(self):
        """El archivo trae código, clave y valor, y no al producto oculto."""
        respuesta = self._descargar(self.usuario)
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], _CONTENT_TYPE_XLSX)
        self.assertIn(
            'Productos_Almacen_Activos_',
            respuesta['Content-Disposition'],
        )
        self.assertIn('.xlsx', respuesta['Content-Disposition'])

        codigos = self._codigos(respuesta)
        self.assertEqual(codigos, ['CON-SAT', 'SIN-SAT'])

        libro = load_workbook(BytesIO(respuesta.content))
        hoja = libro.active
        self.assertEqual(hoja['A1'].value, 'Código')
        self.assertEqual(hoja['L1'].value, 'Clave SAT')
        # CON-SAT queda primero por nombre: «Con clave» < «Sin clave».
        self.assertEqual(hoja['L2'].value, '43211600')
        self.assertEqual(hoja['M2'].value, 10.5)
        self.assertEqual(hoja['N2'].value, 21)
        self.assertEqual(hoja['K3'].value, 'Agotado')
        self.assertNotIn('INACTIVO', codigos)

    def test_filtro_sin_clave_no_mezcla_al_que_ya_tiene(self):
        """clave_sat=sin baja solo el producto vacío."""
        respuesta = self._descargar(self.usuario, {'clave_sat': 'sin'})
        self.assertEqual(self._codigos(respuesta), ['SIN-SAT'])

    def test_la_pagina_de_la_lista_no_recorta_el_archivo(self):
        """page viaja en la URL del botón, pero el Excel trae el catálogo completo."""
        respuesta = self._descargar(self.usuario, {'page': '99'})
        self.assertEqual(self._codigos(respuesta), ['CON-SAT', 'SIN-SAT'])

    def test_sin_permiso_no_recibe_el_archivo(self):
        """Quien no puede ver productos cae en acceso denegado."""
        respuesta = self._descargar(self.sin_permiso)
        self.assertEqual(respuesta.status_code, 302)
        self.assertIn('/almacen/acceso-denegado/', respuesta['Location'])
        self.assertNotEqual(respuesta.get('Content-Type'), _CONTENT_TYPE_XLSX)

    def test_el_boton_arrastra_el_filtro_de_la_pantalla(self):
        """La lista enlaza el Excel con la clave SAT que el usuario eligió."""
        request = self.factory.get('/almacen/productos/', {'clave_sat': 'sin'})
        request.user = self.usuario
        respuesta = lista_productos(request)
        self.assertContains(
            respuesta,
            '/almacen/productos/excel/?clave_sat=sin',
        )
